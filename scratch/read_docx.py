from docx import Document
from pathlib import Path
p=Path(r"C:\Users\ACER\Downloads\Aplikasi Geoverse.docx")
doc=Document(p)
print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")
for i,para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if t:
        print(f"P{i} [{para.style.name}]: {t}")
for ti,table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri,row in enumerate(table.rows):
        vals=[c.text.replace("\\n"," | ").strip() for c in row.cells]
        print(f"T{ti}R{ri}: " + " || ".join(vals))
